# -*- coding: utf-8 -*-
"""Sinh báo cáo tiến độ & mức độ hoàn thành demo CVE-2014-0160 (Heartbleed)."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# ---- Base style ----
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)

def h(text, level=1):
    doc.add_heading(text, level=level)

def p(text, bold=False, italic=False):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = bold
    run.italic = italic
    return par

def bullet(text):
    return doc.add_paragraph(text, style='List Bullet')

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ''
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    return t

# ================= TITLE =================
title = doc.add_heading('BÁO CÁO TIẾN ĐỘ VÀ MỨC ĐỘ HOÀN THÀNH', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run('Demo lỗ hổng Heartbleed — CVE-2014-0160\nMôn: An toàn mạng (BTL)')
run.italic = True

doc.add_paragraph('—' * 40)

# ================= I =================
h('I. TỔNG QUAN DỰ ÁN', 1)
p('Mục tiêu: xây dựng môi trường lab tái hiện lỗ hổng Heartbleed (CVE-2014-0160) trong thư viện '
  'OpenSSL 1.0.1f, đồng thời thực hiện demo khai thác để chứng minh kẻ tấn công có thể thu được '
  'username, password, session ID và cookie của người dùng từ bộ nhớ tiến trình TLS của máy chủ.')
p('Phạm vi thực hiện:')
bullet('Máy chủ nạn nhân (Victim): container chạy OpenSSL 1.0.1f làm reverse proxy TLS, phía sau là '
       'ứng dụng web đầy đủ (Spring Boot + MySQL + JWT + session).')
bullet('Máy tấn công (Attacker): script Python gửi HeartbeatRequest độc hại và đọc dữ liệu rò rỉ '
       '(hb_loop.py, heartbleed_test.py).')

# ================= II =================
h('II. TIẾN ĐỘ THỰC HIỆN', 1)
table(
    ['STT', 'Hạng mục công việc', 'Trạng thái', 'Kết quả / Ghi chú'],
    [
        ['1', 'Môi trường victim (Docker Compose)', 'Hoàn thành', 'mysql + springboot-app + heartbleed-server + nginx-proxy'],
        ['2', 'Web đầy đủ backend + frontend', 'Hoàn thành', 'BTL-AnToanMang: login demo/demo123, quản lý tài liệu, JWT'],
        ['3', 'Tạo session & cookie cho demo', 'Hoàn thành', 'HeartbleedDemoController: JSESSIONID + HEARTBLEED_SESSION'],
        ['4', 'Reverse proxy OpenSSL 1.0.1f (lỗi)', 'Hoàn thành (đã sửa)', 'heartbleed_lab_server.c: đọc đầy đủ request, log dữ liệu bắt được'],
        ['5', 'Script khai thác Heartbleed', 'Hoàn thành (đã sửa)', 'hb_loop.py (lặp leak), heartbleed_test.py (kiểm tra vuln)'],
        ['6', 'Routing nginx (port 80 → proxy TLS lỗi)', 'Hoàn thành', 'nginx.conf'],
        ['7', 'Tài liệu hướng dẫn demo', 'Hoàn thành', 'HUONG_DAN_DEMO.md'],
        ['8', 'Chạy & xác minh end-to-end', 'Đang chờ', 'Cần bật Docker daemon để build và xác minh leak thực tế'],
    ],
)

# ================= III =================
h('III. KIẾN TRÚC HỆ THỐNG', 1)
p('Toàn bộ lưu lượng HTTPS của người dùng đi qua reverse proxy chạy OpenSSL 1.0.1f (bản lỗi). '
  'Proxy giải mã TLS, đọc request dạng plaintext rồi chuyển tiếp sang backend Spring Boot. '
  'Nhờ đó username/password/cookie/session nằm trong bộ nhớ tiến trình TLS và bị script đọc trộm.', italic=True)
p('Sơ đồ luồng:')
bullet('Trình duyệt ──http──▶ nginx:80 ──(302)──▶ /cve-heartbleed/')
bullet('nginx ──proxy_pass https──▶ heartbleed-server:443 (OpenSSL 1.0.1f — LỖI Heartbleed)')
bullet('heartbleed-server ──▶ springboot-app:8080 (web đầy đủ + MySQL)')
bullet('Attacker ──hb_loop.py──▶ heartbleed-server:8082 (HeartbeatRequest độc hại)')

# ================= IV =================
h('IV. CƠ CHẾ KHAI THÁC (TÓM TẮT KỸ THUẬT)', 1)
p('1. OpenSSL 1.0.1f duy trì một freelist read-buffer (rbuf) nội bộ: buffer giải mã TLS (~17 KB) '
  'được tái sử dụng cùng một vùng nhớ cho mọi kết nối.')
p('2. Khi người dùng đăng nhập, request plaintext (gồm body JSON chứa username/password và các header '
  'Cookie/Authorization) được ghi vào buffer đó.')
p('3. Kết nối Heartbeat tiếp theo tái sử dụng đúng buffer này. Lỗi CVE-2014-0160 khiến server đọc tràn '
  '~16 KB từ đầu record heartbeat trở đi — phần đuôi còn sót lại chính là request đăng nhập trước đó.')
p('Kết quả: username, password, session ID và cookie bị rò rỉ mà không để lại dấu vết trong access log.')

# ================= V =================
h('V. MỨC ĐỘ HOÀN THÀNH THEO YÊU CẦU ĐỀ BÀI', 1)
table(
    ['STT', 'Yêu cầu đề bài (nhiemvu2ATM.docx)', 'Mức độ', 'Ghi chú'],
    [
        ['1', 'Khảo sát HTTPS và TLS (mục 1)', '90%', 'Có nội dung lý thuyết; bổ sung hoàn thiện vào báo cáo cuối'],
        ['2', 'Sơ đồ TLS Handshake (mục 2)', '90%', 'Đã có mô tả; cần vẽ hình hoàn chỉnh trong báo cáo cuối'],
        ['3', 'TLS Heartbeat — RFC 6520 (mục 3)', '100%', 'Đã nêu đầy đủ khái niệm và cơ chế'],
        ['4', 'CVE-2014-0160 và kịch bản khai thác (mục 4)', '100%', 'Đã cài đặt môi trường + script khai thác'],
        ['5', 'Xác định môi trường OpenSSL Lab (mục 5)', '100%', 'Victim = container OpenSSL 1.0.1f; Attacker = Python script'],
        ['6', 'Demo thu được username/password/session ID/cookie', '95%', 'Đã chuẩn bị đầy đủ mã nguồn; cần xác minh chạy thực tế'],
    ],
)
p('', bold=False)
par = doc.add_paragraph()
r = par.add_run('MỨC ĐỘ HOÀN THÀNH CHUNG: ')
r.bold = True
par.add_run('khoảng 95% (hạ tầng + mã nguồn + tài liệu). Phần còn lại (≈5%) là bước chạy xác minh '
            'end-to-end trên Docker và hoàn thiện phần lý thuyết + hình vẽ trong báo cáo tổng kết.')

# ================= VI =================
h('VI. CÁC VẤN ĐỀ ĐÃ PHÁT HIỆN VÀ KHẮC PHỤC', 1)
table(
    ['STT', 'Vấn đề', 'Khắc phục'],
    [
        ['1', 'Heap-spray (~128 MB) không được leak vì over-read nằm trọn trong rbuf 17 KB', 'Loại bỏ spray; dùng đúng cơ chế freelist rbuf của OpenSSL'],
        ['2', 'hb_loop.py không nhận ra username/password dạng JSON ("username":"demo")', 'Bổ sung marker "username", "password"'],
        ['3', 'C server chỉ SSL_read 1 lần → thiếu body POST (không leak password, backend trả 400)', 'Đọc đầy đủ request theo Content-Length'],
        ['4', 'Thiếu #include <strings.h> cho strncasecmp', 'Bổ sung header'],
        ['5', 'Thư mục build/web (web tĩnh) không được serve, login giả', 'Xác định dùng web đầy đủ BTL-AnToanMang cho demo'],
    ],
)

# ================= VII =================
h('VII. HƯỚNG DẪN DEMO & CÔNG VIỆC CÒN LẠI', 1)
p('Các bước demo:', bold=True)
bullet('Khởi động Docker Desktop (đảm bảo daemon running).')
bullet('cd E:\\lab2-heartbleed  →  docker compose up --build -d  (build lần đầu ~5–10 phút do biên dịch OpenSSL).')
bullet('Mở terminal riêng: python build\\hb_loop.py  (gửi HeartbeatRequest liên tục).')
bullet('Mở http://localhost/cve-heartbleed/ và đăng nhập: demo / demo123.')
bullet('Quan sát hb_loop.py in ra "username":"demo","password":"demo123" và Cookie: HEARTBLEED_SESSION=...; JSESSIONID=...')
p('Công việc còn lại:')
bullet('Chạy xác minh end-to-end và chụp lại kết quả rò rỉ làm minh chứng.')
bullet('Hoàn thiện phần lý thuyết (HTTPS/TLS, Handshake, Heartbeat, CVE) + hình vẽ vào báo cáo tổng kết.')

# ================= VIII =================
h('VIII. KẾT LUẬN', 1)
p('Lab Heartbleed đã được hoàn thiện về cấu trúc và mã nguồn: web đầy đủ có đăng nhập thật, '
  'session/cookie thật, reverse proxy OpenSSL 1.0.1f lỗi, cùng script khai thác có khả năng thu được '
  'username, password, session ID và cookie. Các vấn đề kỹ thuật quan trọng đã được phát hiện và khắc phục. '
  'Chỉ còn bước chạy xác minh thực tế trên Docker để hoàn tất demo.')

doc.save(r'E:\lab2-heartbleed\BaoCao_TienDo_Heartbleed.docx')
print('OK -> BaoCao_TienDo_Heartbleed.docx')
